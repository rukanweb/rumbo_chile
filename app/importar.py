# -*- coding: utf-8 -*-
"""
importar.py  ·  Trae movimientos de golpe
=========================================
Tres caminos, todos con vista previa antes de guardar nada:
  1. El CSV «Plusvalías y minusvalías» de cada fondo de MyInvestor.
  2. La plantilla genérica (Excel o CSV), rellena a mano o por una IA.
  3. El texto que devuelve la IA, pegado tal cual.

El proceso tiene dos pasos: preparar() lo lee, reconoce los productos y deja un
plan; aplicar() ejecuta ese plan sobre la cartera. La vista previa es aplicar()
sobre una copia, así que lo que ves es exactamente lo que se guardará.
"""

import copy
import csv
import datetime as dt
import io
import re
import unicodedata

from . import almacen, buscar, motor
from .motor import num_es, valor_en

COLUMNAS = ["fecha", "identificador", "nombre", "tipo_producto", "tipo_movimiento",
            "unidades", "importe", "moneda", "comision", "nota"]

# Nombres alternativos que se aceptan en la cabecera.
SINONIMOS = {
    "fecha": ["fecha", "date", "fecha operacion", "fecha valor", "fecha ejecucion"],
    "identificador": ["identificador", "isin", "ticker", "simbolo", "codigo", "id"],
    "nombre": ["nombre", "producto", "descripcion", "name"],
    "tipo_producto": ["tipo producto", "tipo_producto", "clase", "tipo de producto"],
    "tipo_movimiento": ["tipo movimiento", "tipo_movimiento", "movimiento", "operacion", "tipo"],
    "unidades": ["unidades", "participaciones", "titulos", "acciones", "cantidad", "units"],
    "importe": ["importe", "importe eur", "total", "amount", "importe total"],
    "moneda": ["moneda", "divisa", "currency"],
    "comision": ["comision", "comisiones", "fee", "gastos"],
    "nota": ["nota", "notas", "comentario", "observaciones"],
}
TIPO_MOV = {
    "compra": "compra", "suscripcion": "compra", "buy": "compra", "aportacion": "compra",
    "venta": "venta", "reembolso": "venta", "sell": "venta",
    "dividendo": "dividendo", "cupon": "dividendo", "dividend": "dividendo", "interes": "dividendo",
    "comision": "comision", "fee": "comision", "custodia": "comision",
    "saldo": "saldo", "valor": "saldo", "valoracion": "saldo",
}
TIPO_PROD = {
    "fondo": "fondo", "fondo de inversion": "fondo", "etf": "etf", "etp": "etf",
    "accion": "accion", "acciones": "accion", "cripto": "cripto", "criptomoneda": "cripto",
    "commodity": "commodity", "materia prima": "commodity", "oro": "commodity",
    "bono": "bono", "renta fija": "bono", "pension": "pension", "plan de pensiones": "pension",
    "efectivo": "efectivo", "cuenta": "efectivo", "cuenta corriente": "efectivo",
    "inmueble": "inmueble", "piso": "inmueble", "vivienda": "inmueble",
    "deuda": "deuda", "prestamo": "deuda", "hipoteca": "deuda", "otro": "otro",
}


def sin_tildes(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t.replace("_", " ")).strip().lower()


def decodifica(crudo):
    """Bytes de un archivo de texto -> str, probando las codificaciones habituales."""
    if isinstance(crudo, str):
        return crudo
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return crudo.decode(enc)
        except UnicodeDecodeError:
            continue
    return crudo.decode("utf-8", "replace")


# ---------------------------------------------------------------- MyInvestor

CAB_MYINVESTOR = {"fecha": ["fecha fiscal", "fecha"],
                  "coste": ["inversion", "inversión", "coste"],
                  "valor": ["valor de mercado", "valor mercado", "valor"]}


def leer_csv_myinvestor(texto):
    """
    Lee el CSV "Plusvalías y minusvalías" de un fondo de MyInvestor.
    Devuelve (lotes, reembolsos): lotes = [[fecha, coste, valor]] de lo que sigues
    teniendo; reembolsos = [[fecha, resultado]] de lo ya vendido, del que el extracto
    solo da la plusvalía.
    """
    texto = decodifica(texto)
    delim = ";" if texto.count(";") >= texto.count(",") else ","
    filas = [f for f in csv.reader(io.StringIO(texto), delimiter=delim)
             if any((c or "").strip() for c in f)]
    if not filas:
        return [], []

    cab = [(c or "").strip().lower() for c in filas[0]]

    def col(clave, defecto):
        for i, c in enumerate(cab):
            if any(a in c for a in CAB_MYINVESTOR[clave]):
                return i
        return defecto

    ic, ico, iv = col("fecha", 0), col("coste", 1), col("valor", 2)
    lotes, reembolsos = [], []
    for fila in filas[1:]:
        if len(fila) <= max(ic, ico, iv):
            continue
        m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", fila[ic])
        if not m:
            continue
        fecha = f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
        coste, valor = num_es(fila[ico]), num_es(fila[iv])
        if coste <= 0 and valor <= 0:
            # Participaciones ya reembolsadas: solo queda el resultado fiscal.
            if len(fila) > 3 and num_es(fila[3]):
                reembolsos.append([fecha, round(num_es(fila[3]), 2)])
            continue
        lotes.append([fecha, round(coste, 2), round(valor, 2)])
    lotes.sort(key=lambda x: x[0])
    return lotes, reembolsos


def myinvestor_a_movimientos(producto_id, lotes, reembolsos, serie_vl):
    """
    Convierte los lotes en compras. El extracto no trae participaciones, así que
    se calculan como coste / valor liquidativo del día de compra. No se usa el
    valor de mercado del extracto porque MyInvestor lo calcula con un VL de uno o
    dos días antes y saldrían participaciones erróneas.
    """
    # VL con el que MyInvestor valoró el extracto, para los lotes sin VL de compra.
    estimados = sorted(serie_vl[f] * v / c for f, c, v in lotes if serie_vl.get(f) and c > 0)
    vl_extracto = estimados[len(estimados) // 2] if estimados else None
    movs, sin_vl = [], 0
    for fecha, coste, valor in lotes:
        vc = serie_vl.get(fecha)
        if vc:
            unidades = coste / vc
        elif vl_extracto:
            unidades, sin_vl = valor / vl_extracto, sin_vl + 1
        else:
            unidades, sin_vl = 0.0, sin_vl + 1
        movs.append({"fecha": fecha, "producto": producto_id, "tipo": "compra",
                     "unidades": round(unidades, 6), "importe": coste, "nota": "MyInvestor"})
    for fecha, resultado in reembolsos:
        # Del reembolso solo se conoce la plusvalía: se anota como una venta de 0
        # participaciones que cobra ese resultado.
        movs.append({"fecha": fecha, "producto": producto_id, "tipo": "venta",
                     "unidades": 0, "importe": resultado,
                     "nota": "Plusvalía de un reembolso (MyInvestor no da la fecha de venta)"})
    return movs, sin_vl


# ---------------------------------------------------------------- lectura de tablas

def leer_tabla(nombre, contenido):
    """Excel (.xlsx), CSV o texto pegado -> [(número de fila, {columna: valor})]."""
    if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        hoja = libro["Movimientos"] if "Movimientos" in libro.sheetnames else libro.worksheets[0]
        filas = [list(f) for f in hoja.iter_rows(values_only=True)]
    else:
        texto = decodifica(contenido)
        # Lo que devuelve una IA suele venir dentro de un bloque ``` ... ```.
        texto = re.sub(r"^\s*```[a-zA-Z]*\s*$", "", texto, flags=re.M).strip()
        lineas = [l for l in texto.splitlines() if l.strip()]
        if not lineas:
            return [], "El texto está vacío."
        cab = lineas[0]
        delim = max([";", "\t", ","], key=cab.count)
        filas = list(csv.reader(io.StringIO("\n".join(lineas)), delimiter=delim))
    # Se guarda el número de fila real (el que ves en Excel) aunque haya filas en blanco.
    filas = [(n, f) for n, f in enumerate(filas, start=1)
             if f and any(c not in (None, "") and str(c).strip() for c in f)]
    if not filas:
        return [], "No hay ninguna fila con datos."

    cab = [sin_tildes(c) for c in filas[0][1]]
    idx = {}
    for col, alias in SINONIMOS.items():
        for i, c in enumerate(cab):
            if c in alias and i not in idx.values():
                idx[col] = i
                break
    faltan = [c for c in ("fecha", "tipo_movimiento", "importe") if c not in idx]
    if faltan:
        return [], ("No encuentro las columnas " + ", ".join(faltan) + ". La primera fila tiene que ser la "
                    "cabecera de la plantilla: " + ";".join(COLUMNAS))
    if len(filas) < 2:
        if nombre and nombre.lower().endswith((".xlsx", ".xlsm")):
            return [], ("La hoja «Movimientos» está vacía: solo tiene la cabecera. Las filas de la hoja "
                        "«Ejemplo» son de muestra y no se importan; escribe tus operaciones en «Movimientos», "
                        "debajo de la cabecera, guarda el archivo y vuelve a subirlo.")
        return [], "Solo hay cabecera: escribe tus operaciones debajo, una por fila."
    salida = []
    for n, f in filas[1:]:
        salida.append((n, {col: (f[i] if i < len(f) else None) for col, i in idx.items()}))
    return salida, None


def lee_fecha(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    t = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        a, me, di = int(m.group(1)), int(m.group(2)), int(m.group(3))
    else:
        m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", t)
        if not m:
            return None
        a = int(m.group(3)) + (2000 if len(m.group(3)) == 2 else 0)
        me, di = int(m.group(2)), int(m.group(1))
    try:
        return dt.date(a, me, di).isoformat()
    except ValueError:   # un 31 de febrero, un mes 13...
        return None


def lee_numero(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return abs(float(v))
    t = str(v).strip().replace("€", "").replace("$", "").replace(" ", "")
    if not re.fullmatch(r"[-+]?[\d.,]+", t):
        return "error"
    return abs(num_es(t))


# ---------------------------------------------------------------- preparar el plan

class Plan:
    """Lo que se va a importar: productos nuevos, movimientos, saldos y errores."""

    def __init__(self):
        self.productos_nuevos = []   # [{"ref": "nuevo:1", "datos": {...}, "precio": ..., "fecha": ...}]
        self.movimientos = []        # [{"fila", "producto" (id o ref), campos del movimiento, "marcas"}]
        self.valoraciones = []
        self.reemplazar = []         # [(producto id o ref, origen)]: se borran sus importados previos
        self.errores = []            # [{"fila", "mensaje"}]
        self.avisos = []

    def error(self, fila, mensaje):
        self.errores.append({"fila": fila, "mensaje": mensaje})


def _producto_existente(cfg, identificador, nombre):
    ident = (identificador or "").strip().upper()
    for p in cfg.get("productos", []):
        if ident and ident in ((p.get("identificador") or "").upper(), (p.get("codigo") or "").upper()):
            return p
    nom = sin_tildes(nombre)
    if nom and not ident:
        for p in cfg.get("productos", []):
            if nom in (sin_tildes(p.get("nombre")), sin_tildes(p.get("corto"))):
                return p
    return None


class Resolutor:
    """Encuentra (o propone crear) el producto de cada fila, una sola vez por producto."""

    def __init__(self, cfg, plan, carpeta):
        self.cfg, self.plan, self.carpeta = cfg, plan, carpeta
        self.cache, self.series, self.frescas = {}, {}, set()

    def producto(self, fila, identificador, nombre, tipo_prod):
        clave = ((identificador or "").strip().upper(), sin_tildes(nombre) if not identificador else "")
        if clave in self.cache:
            return self.cache[clave]
        ref = None
        existente = _producto_existente(self.cfg, identificador, nombre)
        if existente:
            ref = existente["id"]
        elif identificador:
            n_fallos = len(buscar.FALLOS)
            res = buscar.buscar(identificador)
            if not res and buscar.hubo_fallos_desde(n_fallos):
                self.plan.error(fila, f"No he podido consultar «{identificador}» en internet ahora mismo "
                                      "(la conexión ha fallado o Yahoo/Morningstar no han respondido). "
                                      "Comprueba tu conexión y vuelve a pulsar «Previsualizar»: suele bastar "
                                      "con intentarlo otra vez.")
                ref = None
            elif res:
                r = res[0]
                fi = r.get("ficha") or {}
                datos = {"nombre": nombre or r.get("nombre") or r["codigo"], "tipo": r["tipo"],
                         "identificador": identificador.strip(), "fuente": r["fuente"], "codigo": r["codigo"],
                         "moneda": r.get("moneda") or motor.BASE, "vivo": r.get("vivo") or "",
                         "ter": fi.get("ter"), "riesgo": fi.get("riesgo"), "clase": fi.get("clase") or "",
                         "gestora": fi.get("gestora") or "", "largoPlazo": True}
                ref = self._nuevo(datos, r)
            else:
                self.plan.error(fila, f"No encuentro «{identificador}» con precio en internet. Si no cotiza, "
                                      "créalo en Productos con «Valor anotado a mano» y ese mismo "
                                      "identificador, y vuelve a importar.")
        elif nombre:
            tipo = tipo_prod if tipo_prod in motor.TIPOS else "otro"
            ref = self._nuevo({"nombre": nombre, "tipo": tipo, "fuente": "manual",
                               "largoPlazo": tipo not in ("efectivo", "deuda")}, None)
        else:
            self.plan.error(fila, "No tiene identificador ni nombre: no sé de qué producto es.")
        self.cache[clave] = ref
        return ref

    def _nuevo(self, datos, resultado):
        ref = f"nuevo:{len(self.plan.productos_nuevos) + 1}"
        self.plan.productos_nuevos.append({
            "ref": ref, "datos": datos,
            "precio": resultado and resultado.get("precio"), "monedaPrecio": resultado and resultado.get("moneda"),
            "fecha": resultado and resultado.get("fecha"), "mercado": resultado and resultado.get("mercado")})
        return ref

    def datos(self, ref):
        if ref and ref.startswith("nuevo:"):
            return next(n["datos"] for n in self.plan.productos_nuevos if n["ref"] == ref)
        return almacen.producto(self.cfg, ref)

    def serie(self, ref, fecha=None):
        """Precios en euros del producto; si no llegan a 'fecha', se bajan de nuevo."""
        if ref not in self.series:
            p = self.datos(ref)
            self.series[ref] = motor.serie_producto(p, self.carpeta) if p and p.get("fuente") != "manual" else {}
        s = self.series[ref]
        if fecha and s and max(s) < fecha and ref not in self.frescas:
            self.frescas.add(ref)
            self.series[ref] = s = motor.serie_producto(self.datos(ref), self.carpeta, fresca=True)
        return s


def preparar_tabla(cfg, filas, carpeta):
    """Plantilla o texto de la IA -> Plan."""
    plan = Plan()
    res = Resolutor(cfg, plan, carpeta)
    cambios = {}
    for n, f in filas:
        fecha = lee_fecha(f.get("fecha"))
        tipo = TIPO_MOV.get(sin_tildes(f.get("tipo_movimiento")))
        ident = str(f.get("identificador") or "").strip()
        nombre = str(f.get("nombre") or "").strip()
        tprod = TIPO_PROD.get(sin_tildes(f.get("tipo_producto")), "")
        importe, unidades, comision = (lee_numero(f.get(k)) for k in ("importe", "unidades", "comision"))
        moneda = (str(f.get("moneda") or "").strip() or motor.BASE).upper()
        nota = str(f.get("nota") or "").strip()

        if not fecha:
            plan.error(n, f"La fecha «{f.get('fecha') or ''}» no es válida. Usa el formato 2025-03-10 o 10/03/2025.")
            continue
        if not tipo:
            plan.error(n, f"El tipo de movimiento «{f.get('tipo_movimiento') or ''}» no es válido: "
                          "usa compra, venta, dividendo, comision o saldo.")
            continue
        if "error" in (importe, unidades, comision):
            plan.error(n, "Algún número no se entiende (importe, unidades o comisión).")
            continue
        if not importe and tipo != "saldo":
            plan.error(n, "No tiene importe.")
            continue
        if not ident and not nombre:
            plan.error(n, "No tiene identificador: pon el ISIN o el ticker (o, si es una cuenta, su nombre).")
            continue
        if tipo == "saldo" and not ident:
            tprod = tprod or "efectivo"
        ref = res.producto(n, ident, nombre, tprod)
        if not ref:
            continue
        marcas = []

        # Importes en otra moneda: se pasan a euros con el cambio de ese día.
        if moneda != motor.BASE and importe is not None:
            if moneda not in cambios:
                cambios[moneda] = motor.serie_cambio(moneda, carpeta)
            fx = valor_en(cambios[moneda], fecha, margen=6)
            if not fx:
                plan.error(n, f"No encuentro el cambio de {moneda} a euros del {almacen.fmt_fecha(fecha)}.")
                continue
            importe = importe * fx
            comision = comision * fx if comision else comision
            marcas.append(f"convertido de {moneda}")

        if tipo == "saldo":
            plan.valoraciones.append({"fila": n, "producto": ref, "fecha": fecha, "valor": round(importe or 0, 2)})
            continue

        p = res.datos(ref)
        cotiza = p and p.get("fuente") != "manual" and p.get("tipo") not in almacen.SOLO_SALDO
        if tipo in ("compra", "venta") and not unidades and cotiza:
            # Muchos extractos no dan las participaciones: se calculan con el precio de ese día.
            precio = valor_en(res.serie(ref, fecha), fecha, margen=6)
            if not precio:
                plan.error(n, "No trae unidades y no encuentro el precio de ese día para calcularlas.")
                continue
            neto = importe - (comision or 0) if tipo == "compra" else importe + (comision or 0)
            unidades = neto / precio
            marcas.append("unidades calculadas")
        plan.movimientos.append({"fila": n, "producto": ref, "fecha": fecha, "tipo": tipo,
                                 "unidades": unidades, "importe": round(importe, 2),
                                 "comision": round(comision, 2) if comision else None,
                                 "nota": nota, "marcas": marcas, "origen": "importado"})
    return plan


def preparar_myinvestor(cfg, archivos, carpeta):
    """[(nombre de archivo, bytes)] de MyInvestor -> Plan. Cada archivo sustituye lo que se
    importó antes de ese fondo, porque el extracto siempre trae la foto completa."""
    plan = Plan()
    res = Resolutor(cfg, plan, carpeta)
    for nombre, contenido in archivos:
        m = re.search(r"([A-Z]{2}[A-Z0-9]{9}\d)", (nombre or "").upper())
        if not m:
            plan.error(nombre, "No encuentro el ISIN en el nombre del archivo. Descárgalo otra vez de "
                               "MyInvestor sin cambiarle el nombre (lleva el ISIN del fondo).")
            continue
        isin = m.group(1)
        lotes, reembolsos = leer_csv_myinvestor(contenido)
        if not lotes and not reembolsos:
            plan.error(nombre, "No parece un extracto «Plusvalías y minusvalías» de MyInvestor, o está vacío.")
            continue
        ref = res.producto(nombre, isin, "", "fondo")
        if not ref:
            continue
        p = res.datos(ref)
        if p.get("fuente") == "manual":
            plan.error(nombre, f"«{p.get('corto') or p['nombre']}» se valora a mano: el extracto de MyInvestor "
                               "necesita un fondo con precio en internet.")
            continue
        serie = res.serie(ref, lotes[-1][0] if lotes else None)
        movs, sin_vl = myinvestor_a_movimientos(ref, lotes, reembolsos, serie)
        if sin_vl:
            plan.avisos.append(f"{isin}: {sin_vl} compras sin valor liquidativo de ese día; sus "
                               "participaciones se han estimado con el valor del extracto.")
        plan.reemplazar.append((ref, "myinvestor"))
        for mv in movs:
            plan.movimientos.append({**mv, "fila": nombre, "comision": None, "marcas": [],
                                     "origen": "myinvestor"})
    return plan


# ---------------------------------------------------------------- aplicar

def aplicar(cfg, plan):
    """Ejecuta el plan sobre cfg (la real o una copia) y devuelve el informe."""
    ids, errores = {}, list(plan.errores)
    for n in plan.productos_nuevos:
        prod, _ = almacen.guarda_producto(cfg, dict(n["datos"]))
        ids[n["ref"]] = prod["id"]
    real = lambda ref: ids.get(ref, ref)

    sustituidos = 0
    for ref, origen in plan.reemplazar:
        antes = len(cfg["movimientos"])
        cfg["movimientos"] = [m for m in cfg["movimientos"]
                              if not (m.get("producto") == real(ref) and m.get("origen") == origen)]
        sustituidos += antes - len(cfg["movimientos"])

    def firma(m):
        return (m["producto"], m["fecha"], m["tipo"], round(float(m.get("unidades") or 0), 4),
                round(float(m.get("importe") or 0), 2))
    existentes = {firma(m) for m in cfg["movimientos"]}

    añadidos, duplicados, filas = 0, 0, []
    orden = {"compra": 0, "comision": 1, "dividendo": 2, "venta": 3}
    for mv in sorted(plan.movimientos, key=lambda m: (m["fecha"], orden.get(m["tipo"], 9))):
        datos = {k: mv[k] for k in ("fecha", "tipo", "unidades", "importe", "comision", "nota") if mv.get(k) is not None}
        datos["producto"] = real(mv["producto"])
        estado = "nuevo"
        if firma({**datos, "unidades": datos.get("unidades", 0)}) in existentes:
            duplicados += 1
            estado = "repetido"
        else:
            try:
                guardado = almacen.guarda_movimiento(cfg, datos)
                guardado["origen"] = mv["origen"]
                añadidos += 1
            except almacen.ErrorValidacion as e:
                errores.append({"fila": mv["fila"], "mensaje": " ".join(e.errores)})
                estado = "error"
        filas.append({"fila": mv["fila"], "fecha": mv["fecha"], "producto": datos["producto"], "tipo": mv["tipo"],
                      "unidades": mv.get("unidades"), "importe": mv["importe"], "marcas": mv["marcas"],
                      "estado": estado})

    saldos = 0
    for v in plan.valoraciones:
        try:
            almacen.guarda_valoracion(cfg, {"producto": real(v["producto"]), "fecha": v["fecha"], "valor": v["valor"]})
            saldos += 1
            filas.append({"fila": v["fila"], "fecha": v["fecha"], "producto": real(v["producto"]), "tipo": "saldo",
                          "unidades": None, "importe": v["valor"], "marcas": [], "estado": "nuevo"})
        except almacen.ErrorValidacion as e:
            errores.append({"fila": v["fila"], "mensaje": " ".join(e.errores)})

    def total(t):
        return round(sum(f["importe"] for f in filas if f["tipo"] == t and f["estado"] == "nuevo"), 2)
    nombres = {p["id"]: p.get("corto") or p["nombre"] for p in cfg["productos"]}
    for f in filas:
        f["productoNombre"] = nombres.get(f["producto"], f["producto"])
    return {
        "añadidos": añadidos, "repetidos": duplicados, "saldos": saldos, "sustituidos": sustituidos,
        "errores": sorted(errores, key=lambda e: (str(type(e["fila"])), str(e["fila"]).zfill(6))),
        "avisos": plan.avisos,
        "totales": {"compras": total("compra"), "ventas": total("venta"),
                    "dividendos": total("dividendo"), "comisiones": total("comision"),
                    "numCompras": sum(1 for f in filas if f["tipo"] == "compra" and f["estado"] == "nuevo")},
        "productosNuevos": [{**n["datos"], "id": ids.get(n["ref"]), "precio": n["precio"],
                             "monedaPrecio": n["monedaPrecio"], "fechaPrecio": n["fecha"],
                             "mercado": n["mercado"]} for n in plan.productos_nuevos],
        "filas": sorted(filas, key=lambda f: f["fecha"], reverse=True),
    }


def vista_previa(cfg, plan):
    """El mismo informe que dará aplicar(), sin tocar la cartera de verdad."""
    return aplicar(copy.deepcopy(cfg), plan)
